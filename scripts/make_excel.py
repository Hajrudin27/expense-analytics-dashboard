"""
Pack the six CSVs into ONE Excel workbook, each table on its own sheet and
formatted as a real Excel Table (ListObject).

Why: Power BI Service's CSV upload creates a separate semantic model per
file, which makes relationships impossible. Its Excel upload reads every
ListObject in the workbook into a single model — which is what we need.
"""

import csv
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = os.path.join(os.path.dirname(__file__), "..")
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "powerbi", "ExpenseAnalytics.xlsx")

FONT = "Arial"

DATE_COLS = {"expense_date", "settlement_date", "date_key", "joined_date"}
NUM_COLS = {"amount_dkk", "share_amount_dkk"}
INT_COLS = {
    "person_id", "category_id", "expense_id", "share_id", "settlement_id",
    "paid_by_person_id", "from_person_id", "to_person_id", "year",
    "month_number", "day_of_month", "participant_count",
}
BOOL_COLS = {"is_weekend", "is_payer"}

SHEETS = [
    ("dim_person", "dim_person.csv"),
    ("dim_category", "dim_category.csv"),
    ("dim_date", "dim_date.csv"),
    ("fact_expense", "fact_expense.csv"),
    ("fact_expense_share", "fact_expense_share.csv"),
    ("fact_settlement", "fact_settlement.csv"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)


def convert(col, raw):
    if raw == "":
        return None
    if col in DATE_COLS:
        return datetime.strptime(raw, "%Y-%m-%d").date()
    if col in INT_COLS:
        return int(raw)
    if col in NUM_COLS:
        return float(raw)
    if col in BOOL_COLS:
        return raw.lower() == "true"
    return raw


def add_sheet(wb, sheet_name, filename, first):
    rows = list(csv.DictReader(open(os.path.join(DATA, filename), encoding="utf-8")))
    cols = list(rows[0].keys())

    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_name

    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    for r in rows:
        ws.append([convert(c, r[c]) for c in cols])

    # Number / date formats and body font
    for c_idx, col in enumerate(cols, start=1):
        letter = get_column_letter(c_idx)
        if col in DATE_COLS:
            fmt = "yyyy-mm-dd"
        elif col in NUM_COLS:
            fmt = "#,##0.00"
        else:
            fmt = None
        for r_idx in range(2, len(rows) + 2):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = BODY_FONT
            if fmt:
                cell.number_format = fmt

        width = max(len(col) + 4, min(28, max((len(str(r[col])) for r in rows), default=8) + 3))
        ws.column_dimensions[letter].width = width

    # The ListObject — this is what Power BI actually detects.
    ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    table = Table(displayName=sheet_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    ws.freeze_panes = "A2"

    print(f"  {sheet_name:22s} {len(rows):>5,} rows x {len(cols)} cols   ({ref})")
    return len(rows)


def main():
    print("Building ExpenseAnalytics.xlsx ...")
    wb = Workbook()
    total = 0
    for i, (sheet, filename) in enumerate(SHEETS):
        total += add_sheet(wb, sheet, filename, first=(i == 0))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"\n  {len(SHEETS)} tables, {total:,} data rows")
    print(f"  wrote {OUT} ({os.path.getsize(OUT):,} bytes)")


if __name__ == "__main__":
    main()
